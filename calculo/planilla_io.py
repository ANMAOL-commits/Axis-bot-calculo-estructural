"""
Lee los datos de entrada de la 'Planilla Unificada de Cálculo Estructural' (.xlsx) y arma el diccionario de proyecto.
Se leen los VALORES guardados en la planilla (abrir y guardar en Excel una vez si se editó con otro programa).
"""
import warnings

warnings.filterwarnings("ignore")
from openpyxl import load_workbook

from .tablas import F1, VIGUETAS


def _v(ws, ref):
    v = ws[ref].value
    return v.strip() if isinstance(v, str) else v


def _fila(ws, empieza, col="B", hasta=150):
    """Busca la fila cuyo texto (en la columna indicada) empieza con 'empieza'. Así el lector no depende de filas fijas."""
    for r in range(1, hasta + 1):
        v = ws[f"{col}{r}"].value
        if isinstance(v, str) and v.strip().startswith(empieza):
            return r
    raise ValueError(f"No se encontró '{empieza}' en la hoja '{ws.title}': la planilla fue modificada (se movieron filas o se cambió un rótulo).")


def _clave(valor, dic):
    """Devuelve la clave de 'dic' equivalente a 'valor' (tolera 'Serie 3a' / '3a' y etiquetas largas o cortas)."""
    if valor is None:
        return valor
    s = str(valor).strip()
    if s in dic:
        return s
    corto = s.replace("Serie ", "")
    for k in dic:
        if k.replace("Serie ", "") == corto or k.startswith(s[:14]) or s.startswith(k[:14]):
            return k
    return s


def leer_planilla(ruta):
    wb = load_workbook(ruta, data_only=True)
    req = ["Datos", "Cargas", "Losas", "Vigas", "Columnas", "Sismo 1 Accion", "Sismo 2 Distribucion", "Muros", "Fundaciones", "Voladizos", "Viguetas y Correas"]
    faltan = [h for h in req if h not in wb.sheetnames]
    if faltan:
        raise ValueError("El archivo no es la Planilla Unificada: faltan las hojas " + ", ".join(faltan))
    avisos = []
    D, Ca, Lo, Vi, Co, S1, S2, Mu, Fu, Vo, VC = (wb[h] for h in req)
    P = {"avisos": avisos}
    # ---- datos generales
    P["datos"] = dict(obra=_v(D, "C5"), ubicacion=_v(D, "C6"), comitente=_v(D, "C7"), proyectista=_v(D, "C8"), fecha=str(_v(D, "C9") or ""), revision=str(_v(D, "C10") or ""),
                      clase_H=_v(D, "C13"), gHA=_v(D, "C17"), gHS=_v(D, "C18"), cc_vc=_v(D, "C19"), cc_l=_v(D, "C20"), cc_z=_v(D, "C21"), acero_long=_v(D, "C24"), acero_est=_v(D, "C26"),
                      Es=_v(D, "C28"), s_adm=_v(D, "C31"), Df=_v(D, "C33"), g_suelo=_v(D, "C34"), g_mamp=_v(D, "C35"), mamp_tipo=_v(D, "C38"), mortero=_v(D, "C39"))
    # ---- cargas
    def capas(r1, r2):
        out = []
        for r in range(r1, r2 + 1):
            if _v(Ca, f"B{r}"):
                out.append(dict(desc=_v(Ca, f"B{r}"), g=_v(Ca, f"C{r}"), q=_v(Ca, f"D{r}"), e=_v(Ca, f"E{r}"), b=_v(Ca, f"F{r}"), sep=_v(Ca, f"G{r}"), alpha=_v(Ca, f"H{r}")))
        return out
    # filas ubicadas por rótulo (funciona con la planilla v1.0 y v1.1)
    rT = _fila(Ca, "D TECHO / CUBIERTA")
    rE = _fila(Ca, "D ENTREPISO")
    rLT = _fila(Ca, "Sobrecarga de uso sobre el techo")      # L (techo); las dos filas anteriores son Lr (techo) y L (entrepiso)
    rLrE = _fila(Ca, "Sobrecarga de techo sobre el entrepiso")
    rfL = _fila(Ca, "Factor de L en las combinaciones")
    rh = _fila(Ca, "Peso propio estimado de la losa")
    rw = _fila(Ca, "Tipo de mampostería (elegir de la lista)") + 1

    def val(r, nombre, defecto=0):
        v = _v(Ca, f"I{r}")
        if v is None:
            avisos.append(f"Cargas: '{nombre}' no tiene valor guardado (abrir y guardar la planilla en Excel); se tomó {defecto}.")
            return defecto
        return v

    cT = dict(capas=capas(rT - 10, rT - 3), correa=dict(desc=_v(Ca, f"B{rT-1}"), kNm=_v(Ca, f"C{rT-1}"), sep=_v(Ca, f"G{rT-1}"), alpha=_v(Ca, f"H{rT-1}")))
    cE = dict(capas=capas(rE - 8, rE - 1))
    P["cargas"] = dict(T=cT, E=cE, Lr_T=val(rLT - 2, "sobrecarga de techo Lr"), L_E=val(rLT - 1, "sobrecarga de entrepiso L"), L_T=val(rLT, "L sobre techo"), Lr_E=val(rLrE, "Lr sobre entrepiso"),
                       fL=val(rfL, "factor fL", 0.5), h_est=val(rh, "espesor estimado de losa", 0.12), desc_Lr_T=_v(Ca, f"B{rLT-2}") or "", desc_L_E=_v(Ca, f"B{rLT-1}") or "",
                       muros=[dict(desc=_v(Ca, f"B{r}"), g=_v(Ca, f"C{r}"), e=_v(Ca, f"E{r}"), h=_v(Ca, f"F{r}")) for r in range(rw, rw + 5) if _v(Ca, f"B{r}")])
    # ---- losas
    P["losas"] = []
    for i in range(6):
        if not _v(Lo, f"B{6+i}"):
            continue
        P["losas"].append(dict(id=_v(Lo, f"B{6+i}"), tipo=_v(Lo, f"C{6+i}"), apoyo=_v(Lo, f"D{6+i}"), L=_v(Lo, f"E{6+i}"), h=_v(Lo, f"I{6+i}"), nivel=_v(Lo, f"C{16+i}"), esquema=_v(Lo, f"D{16+i}"),
                               phi_t=_v(Lo, f"I{26+i}"), s_t=_v(Lo, f"J{26+i}"), phi_a=_v(Lo, f"I{36+i}"), s_a=_v(Lo, f"J{36+i}"), phi_r=_v(Lo, f"G{45+i}"), s_r=_v(Lo, f"H{45+i}")))
    # ---- vigas
    P["vigas"] = []
    for i in range(6):
        if not _v(Vi, f"B{6+i}"):
            continue
        if _v(Vi, f"H{15+i}") is None:
            avisos.append(f"Viga {_v(Vi, f'B{6+i}')}: la reacción de losas no tiene valor guardado (abrir y guardar la planilla en Excel).")
        P["vigas"].append(dict(id=_v(Vi, f"B{6+i}"), cond=_v(Vi, f"C{6+i}"), L=_v(Vi, f"D{6+i}"), h=_v(Vi, f"H{6+i}"), esquema=_v(Vi, f"C{15+i}"), bw=_v(Vi, f"D{15+i}"), phi_e=_v(Vi, f"G{15+i}"),
                               R_losas=_v(Vi, f"H{15+i}"), D_adic=_v(Vi, f"I{15+i}"), L_adic=_v(Vi, f"J{15+i}"),
                               flex_apoyo=dict(n=_v(Vi, f"E{25+2*i}"), phi=_v(Vi, f"F{25+2*i}")), flex_tramo=dict(n=_v(Vi, f"E{26+2*i}"), phi=_v(Vi, f"F{26+2*i}")), ramas=_v(Vi, f"J{41+i}"), sep=_v(Vi, f"K{41+i}")))
    # ---- viguetas y correas
    P["modo_viguetas"] = _v(VC, "E5") or "Servicio"
    P["viguetas"] = [dict(id=_v(VC, f"B{7+i}"), nivel=_v(VC, f"C{7+i}"), serie=_clave(_v(VC, f"D{7+i}"), VIGUETAS), L=_v(VC, f"E{7+i}"), sep=_v(VC, f"F{7+i}"), pp=_v(VC, f"G{7+i}")) for i in range(4) if _v(VC, f"B{7+i}")]
    tb = wb["Tablas"]
    md, mk, rr = [], [], 5
    while tb[f"Q{rr}"].value is not None:           # tabla de maderas (sección y peso por metro)
        md.append(str(tb[f"Q{rr}"].value).strip()); mk.append(tb[f"R{rr}"].value); rr += 1
    P["correas"] = []
    for i in range(4):
        r = 15 + i
        if not _v(VC, f"B{r}"):
            continue
        sec = _v(VC, f"D{r}")
        w = mk[md.index(sec)] if sec in md else 0
        if sec and sec not in md:
            avisos.append(f"Correa {_v(VC, f'B{r}')}: la sección '{sec}' no figura en la tabla de maderas (peso propio = 0).")
        P["correas"].append(dict(id=_v(VC, f"B{r}"), nivel=_v(VC, f"C{r}"), seccion=sec, b=_v(VC, f"E{r}"), h=_v(VC, f"F{r}"), w=w, L=_v(VC, f"G{r}"), sep=_v(VC, f"H{r}"), alpha=_v(VC, f"I{r}"), E=_v(VC, f"J{r}"),
                                 Fb=_v(VC, f"K{r}"), Fv=_v(VC, f"L{r}"), x_flecha=_v(VC, f"Q{r}")))
    # ---- columnas
    P["columnas"] = []
    for i in range(6):
        if not _v(Co, f"B{6+i}"):
            continue
        P["columnas"].append(dict(id=_v(Co, f"B{6+i}"), b=_v(Co, f"C{6+i}"), h=_v(Co, f"D{6+i}"), lu=_v(Co, f"E{6+i}"), A_T=_v(Co, f"F{6+i}"), A_E=_v(Co, f"G{6+i}"), R_vigas=_v(Co, f"H{6+i}"), Pu_sup=_v(Co, f"I{6+i}"),
                                  Mu=_v(Co, f"L{6+i}"), k=_v(Co, f"M{6+i}"), arriostrado=_v(Co, f"N{6+i}") or "SI", m1m2=_v(Co, f"O{6+i}"),
                                  n_tot=_v(Co, f"C{16+i}"), n_cara=_v(Co, f"D{16+i}"), phi=_v(Co, f"E{16+i}"), phi_e=_v(Co, f"F{16+i}"), s_est=_v(Co, f"G{16+i}")))
    # ---- sismo 1
    P["sismo1"] = dict(zona=_v(S1, "C5"), tipo_esp=_v(S1, "C6"), Na=_v(S1, "C7"), Nv=_v(S1, "C8"), grupo=_v(S1, "C9"), sistema=_v(S1, "C11"), A=_v(S1, "C12"), Ar=_v(S1, "C13"), H=_v(S1, "C17"), sist_T=_v(S1, "C18"), T_calc=_v(S1, "C19"),
                       niveles=[dict(nombre=_v(S1, f"B{44+i}"), hk=_v(S1, f"C{44+i}"), A=_v(S1, f"D{44+i}"), tipo=_v(S1, f"E{44+i}"), h_losa=_v(S1, f"F{44+i}"), destino=_clave(_v(S1, f"G{44+i}"), F1), W_otros=_v(S1, f"M{44+i}")) for i in range(3)],
                       muros=[dict(desc=_v(S1, f"B{54+i}"), piso=_v(S1, f"C{54+i}"), long=_v(S1, f"D{54+i}"), esp=_v(S1, f"E{54+i}"), alt=_v(S1, f"F{54+i}")) for i in range(8) if _v(S1, f"B{54+i}")],
                       nivel_sel=_v(S1, "C74"))
    # ---- sismo 2
    els = []
    for i in range(14):
        r = 27 + i
        if _v(S2, f"B{r}"):
            els.append(dict(id=_v(S2, f"B{r}"), dir=_v(S2, f"C{r}"), tipo=_v(S2, f"D{r}"), e=_v(S2, f"E{r}"), L=_v(S2, f"F{r}"), h=_v(S2, f"G{r}"), pos=_v(S2, f"H{r}"), R_manual=_v(S2, f"I{r}")))
    P["sismo2"] = dict(irregularidad=_v(S2, "C6"), Lx=_v(S2, "C8"), Ly=_v(S2, "C9"), cm=[dict(desc=_v(S2, f"B{14+i}"), A=_v(S2, f"C{14+i}"), x=_v(S2, f"E{14+i}"), y=_v(S2, f"F{14+i}")) for i in range(8) if _v(S2, f"C{14+i}")], elementos=els)
    # ---- muros (datos por elemento)
    ND, tab = {}, {}
    for i in range(14):
        idr = _v(S2, f"B{27+i}")
        if idr and _v(Mu, f"G{6+i}") is not None:
            ND[idr] = _v(Mu, f"G{6+i}")
        if idr and i < 6:
            tab[idr] = dict(Nu=_v(Mu, f"F{24+i}"), phi=_v(Mu, f"I{24+i}"), sep=_v(Mu, f"J{24+i}"), cortinas=_v(Mu, f"K{24+i}"))
    P["muros"] = dict(ND=ND, tabiques=tab)
    # ---- voladizos
    P["voladizos"] = dict(columnas=[dict(id=_v(Vo, f"B{7+i}"), b=_v(Vo, f"C{7+i}"), h=_v(Vo, f"D{7+i}"), H=_v(Vo, f"E{7+i}"), Q1=_v(Vo, f"F{7+i}"), n_barras=_v(Vo, f"K{7+i}"), phi=_v(Vo, f"L{7+i}"), phi_e=_v(Vo, f"M{7+i}")) for i in range(4) if _v(Vo, f"B{7+i}")],
                          componentes=[dict(desc=_v(Vo, f"B{15+i}"), Wp=_v(Vo, f"C{15+i}"), z=_v(Vo, f"D{15+i}"), ap=_v(Vo, f"F{15+i}"), Rp=_v(Vo, f"G{15+i}"), Ip=_v(Vo, f"H{15+i}")) for i in range(4) if _v(Vo, f"B{15+i}")])
    # ---- fundaciones
    cim = [dict(id=_v(Fu, f"B{6+i}"), q1=_v(Fu, f"C{6+i}"), L1=_v(Fu, f"D{6+i}"), q2=_v(Fu, f"E{6+i}"), L2=_v(Fu, f"F{6+i}"), e=_v(Fu, f"G{6+i}"), h=_v(Fu, f"H{6+i}"), B=_v(Fu, f"K{6+i}"), hc=_v(Fu, f"L{6+i}")) for i in range(6) if _v(Fu, f"B{6+i}")]
    bases = []
    for i in range(4):
        if _v(Fu, f"B{16+i}"):
            if _v(Fu, f"E{16+i}") is None:
                avisos.append(f"Base {_v(Fu, f'B{16+i}')}: Pu sin valor guardado (se tomará de la columna correspondiente si existe).")
            bases.append(dict(id=_v(Fu, f"B{16+i}"), bc=_v(Fu, f"C{16+i}"), hc=_v(Fu, f"D{16+i}"), Pu=_v(Fu, f"E{16+i}"), B=_v(Fu, f"H{16+i}"), h=_v(Fu, f"I{16+i}"), phi=_v(Fu, f"J{24+i}"), sep=_v(Fu, f"K{24+i}"), idx=i))
    P["fundaciones"] = dict(cimientos=cim, bases=bases)
    return P
